from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement as SharedOxmlElement
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
OUTPUT_DIR = ROOT / "outputs"
ASSET_DIR = OUTPUT_DIR / "tomato_egg_assets"
DOCX_PATH = OUTPUT_DIR / "西红柿炒鸡蛋的最佳做法.docx"

FONT_UI = "微软雅黑"
FONT_BODY = "仿宋"
FONT_UI_PATH = Path(r"C:\Windows\Fonts\msyh.ttc")
FONT_UI_BOLD_PATH = Path(r"C:\Windows\Fonts\msyhbd.ttc")
FONT_BODY_PATH = Path(r"C:\Windows\Fonts\simfang.ttf")

BODY_TEXT = RGBColor(42, 42, 42)
TITLE_COLOR = RGBColor(176, 58, 46)
HEADING_COLOR = RGBColor(132, 33, 26)
ACCENT_COLOR = RGBColor(198, 86, 64)
MUTED_COLOR = RGBColor(105, 105, 105)
TABLE_BORDER = "C7795D"
TABLE_FILL = "F8E8DF"
PAGE_WIDTH_DXA = 9360


def set_east_asia_font(run, font_name: str, size: int | None = None, *, bold: bool = False, color=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def style_paragraph(paragraph, *, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, before=6, after=6):
    fmt = paragraph.paragraph_format
    fmt.line_spacing = line_spacing
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    paragraph.alignment = alignment


def format_body_paragraph(paragraph):
    style_paragraph(paragraph)
    for run in paragraph.runs:
        set_east_asia_font(run, FONT_BODY, 12, color=BODY_TEXT)


def add_body_paragraph(doc: Document, text: str):
    p = doc.add_paragraph()
    style_paragraph(p)
    run = p.add_run(text)
    set_east_asia_font(run, FONT_BODY, 12, color=BODY_TEXT)
    return p


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph()
    style_paragraph(
        p,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        line_spacing=1.25,
        before=12 if level == 1 else 8,
        after=6,
    )
    run = p.add_run(text)
    size = 16 if level == 1 else 14
    set_east_asia_font(run, FONT_UI, size, bold=True, color=HEADING_COLOR)
    return p


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    style_paragraph(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.15, before=2, after=8)
    run = p.add_run(text)
    set_east_asia_font(run, FONT_BODY, 10, color=MUTED_COLOR)
    return p


def set_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_east_asia_font(run, FONT_UI, 9, color=MUTED_COLOR)
    fld_char_begin = SharedOxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = SharedOxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_end = SharedOxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)
    run2 = paragraph.add_run(" 页")
    set_east_asia_font(run2, FONT_UI, 9, color=MUTED_COLOR)


def add_header_footer(section, header_text: str):
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    header_p = section.header.paragraphs[0]
    header_p.clear()
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_paragraph(header_p, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.0, before=0, after=0)
    run = header_p.add_run(header_text)
    set_east_asia_font(run, FONT_UI, 9, color=MUTED_COLOR)

    p_pr = header_p._element.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "E1D7D1")
    border.append(bottom)
    p_pr.append(border)

    footer_p = section.footer.paragraphs[0]
    footer_p.clear()
    set_page_number(footer_p)


def set_section_layout(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, *, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        elem = tc_mar.find(qn(f"w:{key}"))
        if elem is None:
            elem = OxmlElement(f"w:{key}")
            tc_mar.append(elem)
        elem.set(qn("w:w"), str(value))
        elem.set(qn("w:type"), "dxa")


def set_table_fixed_layout(table, widths_dxa: Iterable[int]):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(PAGE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    existing_borders = tbl_pr.first_child_found_in("w:tblBorders")
    if existing_borders is not None:
        tbl_pr.remove(existing_borders)
    tbl_borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), TABLE_BORDER)
        tbl_borders.append(border)
    tbl_pr.append(tbl_borders)

    grid = tbl.tblGrid
    while len(grid):
        grid.remove(grid[0])
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_table_text(cell, *, bold=False, center=False):
    for paragraph in cell.paragraphs:
        style_paragraph(
            paragraph,
            alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT,
            line_spacing=1.2,
            before=2,
            after=2,
        )
        for run in paragraph.runs:
            set_east_asia_font(run, FONT_BODY, 10 if not bold else 10, bold=bold, color=BODY_TEXT)


def add_recipe_table(doc: Document):
    add_heading(doc, "调味黄金比例", level=2)
    intro = add_body_paragraph(
        doc,
        "真正好吃的西红柿炒鸡蛋，味道不在于调料放得多，而在于比例稳。下面这张表可以当作两到三人份的基准，再按番茄酸甜度和家里口味微调。",
    )
    intro.paragraph_format.keep_with_next = True

    table = doc.add_table(rows=1, cols=4)
    widths = [1700, 2600, 2100, 2960]
    set_table_fixed_layout(table, widths)
    headers = ["项目", "建议用量", "最佳加入时机", "作用说明"]
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = text
        shade_cell(cell, TABLE_FILL)
        style_table_text(cell, bold=True, center=True)

    rows = [
        ("盐", "鸡蛋 1.5 克；番茄 1.5 克", "蛋液调味一次，出锅前校正一次", "分两次放盐更稳，不容易一口淡一口咸。"),
        ("白糖", "3-5 克", "番茄下锅出汁后", "提鲜压酸，让番茄味更圆润，不是为了做甜口。"),
        ("料酒", "3 毫升", "打蛋时加入", "帮助去腥并让蛋香更干净。"),
        ("清水或番茄汁", "15-30 毫升", "番茄炒软后酌情加入", "用来衔接汤汁，让鸡蛋更容易挂味。"),
        ("食用油", "18-22 毫升", "煎蛋 12 毫升，炒番茄 6-10 毫升", "油量足，鸡蛋才蓬；分锅香味也更清楚。"),
    ]
    for row_data in rows:
        row = table.add_row()
        for idx, text in enumerate(row_data):
            row.cells[idx].text = text
            style_table_text(row.cells[idx], center=idx == 1)

    doc.add_paragraph()


def add_problem_table(doc: Document):
    add_heading(doc, "常见问题速查表", level=2)
    table = doc.add_table(rows=1, cols=4)
    widths = [1800, 2500, 2500, 2560]
    set_table_fixed_layout(table, widths)
    headers = ["问题表现", "常见原因", "补救办法", "下次避免重点"]
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = text
        shade_cell(cell, TABLE_FILL)
        style_table_text(cell, bold=True, center=True)

    rows = [
        ("鸡蛋发老、发硬", "火太小、炒太久，或蛋液里盐过多", "立刻关火，借番茄汁回润", "鸡蛋炒到八成熟就盛出，余温会继续熟。"),
        ("番茄不出汁", "番茄太生或下锅后频繁翻动", "加一撮盐并盖锅 20 秒", "选熟软番茄，切块后静置几分钟再炒。"),
        ("整道菜发水", "番茄和鸡蛋一起久煮", "大火收汁 20-30 秒", "先分炒后回锅，汤汁只收不久炖。"),
        ("味道寡淡", "盐、糖都放得太晚或太少", "沿锅边补少量盐糖", "先给蛋液底味，再给番茄定味。"),
    ]
    for row_data in rows:
        row = table.add_row()
        for idx, text in enumerate(row_data):
            row.cells[idx].text = text
            style_table_text(row.cells[idx], center=idx == 0)

    doc.add_paragraph()


def add_step(doc: Document, title: str, body: str):
    p = doc.add_paragraph(style="List Number")
    style_paragraph(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.5, before=6, after=4)
    lead = p.add_run(title)
    set_east_asia_font(lead, FONT_UI, 12, bold=True, color=BODY_TEXT)
    rest = p.add_run("：" + body)
    set_east_asia_font(rest, FONT_BODY, 12, color=BODY_TEXT)
    return p


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.35, before=2, after=2)
    run = p.add_run(text)
    set_east_asia_font(run, FONT_BODY, 12, color=BODY_TEXT)
    return p


def generate_illustration(path: Path, title: str, subtitle: str, palette: tuple[str, str, str], mode: str):
    img = Image.new("RGB", (1200, 700), palette[0])
    draw = ImageDraw.Draw(img)
    title_font = ImageFont.truetype(str(FONT_UI_BOLD_PATH), 48)
    sub_font = ImageFont.truetype(str(FONT_BODY_PATH), 26)
    label_font = ImageFont.truetype(str(FONT_UI_PATH), 24)

    draw.rounded_rectangle((40, 40, 1160, 660), radius=36, outline=palette[1], width=4, fill=palette[0])
    draw.text((70, 72), title, font=title_font, fill=palette[1])
    draw.text((72, 140), subtitle, font=sub_font, fill=palette[2])

    if mode == "ingredients":
        draw.ellipse((120, 240, 410, 530), fill="#D94841", outline="#A62C2B", width=6)
        draw.rectangle((240, 210, 285, 270), fill="#4D8C57")
        draw.ellipse((520, 260, 700, 430), fill="#F6D766", outline="#C99818", width=5)
        draw.ellipse((700, 250, 940, 470), fill="#FFF9EC", outline="#D6C8A1", width=5)
        draw.text((160, 545), "成熟番茄", font=label_font, fill=palette[1])
        draw.text((575, 490), "鸡蛋与底味", font=label_font, fill=palette[1])
    elif mode == "knife":
        draw.polygon([(170, 500), (410, 250), (490, 330), (250, 580)], fill="#D94F45", outline="#9E2B25")
        draw.polygon([(500, 220), (1030, 320), (990, 400), (460, 300)], fill="#C0C6CF", outline="#858C94")
        draw.rectangle((920, 270, 1080, 450), fill="#8B5E3C", outline="#6C4529")
        draw.text((140, 565), "切块保汁，去蒂不去瓤", font=label_font, fill=palette[1])
    elif mode == "heat":
        draw.arc((110, 260, 370, 610), start=180, end=360, fill="#3E3E3E", width=18)
        draw.rectangle((170, 420, 310, 480), fill="#3E3E3E")
        draw.ellipse((190, 310, 920, 570), fill="#22262A", outline="#111", width=6)
        draw.ellipse((290, 360, 820, 520), fill="#F2B84B")
        draw.ellipse((350, 390, 760, 500), fill="#FFE08A")
        for offset in [0, 90, 180]:
            draw.polygon([(460 + offset, 250), (500 + offset, 170), (540 + offset, 250)], fill="#F06A3C")
        draw.text((170, 585), "热锅宽油，蛋液遇热先鼓起再定型", font=label_font, fill=palette[1])
    elif mode == "steps":
        for idx, x in enumerate([150, 390, 630, 870], start=1):
            draw.rounded_rectangle((x, 240, x + 160, 450), radius=26, fill="#FFF5E7", outline=palette[1], width=4)
            draw.ellipse((x + 48, 275, x + 112, 339), fill=palette[1])
            draw.text((x + 68, 288), str(idx), font=label_font, fill="white")
        labels = ["打蛋调味", "番茄出汁", "回锅合炒", "收汁出锅"]
        for idx, x in enumerate([130, 370, 610, 850]):
            draw.text((x, 475), labels[idx], font=label_font, fill=palette[1])
    elif mode == "fix":
        draw.rounded_rectangle((120, 220, 520, 520), radius=32, fill="#FFF4EF", outline="#C7795D", width=5)
        draw.rounded_rectangle((660, 220, 1060, 520), radius=32, fill="#EEF7EF", outline="#6E9F67", width=5)
        draw.text((205, 280), "问题", font=title_font, fill="#B24934")
        draw.text((730, 280), "解决", font=title_font, fill="#4D8051")
        draw.text((190, 390), "蛋老、汤多、味淡", font=label_font, fill="#7A3D2A")
        draw.text((720, 390), "缩短时间、先分炒、分次调味", font=label_font, fill="#406B44")
    elif mode == "nutrition":
        draw.ellipse((210, 205, 960, 585), fill="#F5F5F5", outline="#D6D6D6", width=5)
        draw.pieslice((250, 245, 920, 545), start=0, end=110, fill="#A3D39C")
        draw.pieslice((250, 245, 920, 545), start=110, end=240, fill="#F3D26C")
        draw.pieslice((250, 245, 920, 545), start=240, end=360, fill="#E4745F")
        draw.text((260, 575), "番茄提供维生素，鸡蛋补优质蛋白，主食搭配更均衡", font=label_font, fill=palette[1])
    else:
        draw.ellipse((220, 200, 980, 580), fill="#C83F35", outline="#9B2C22", width=8)
        draw.ellipse((300, 270, 900, 530), fill="#F1C553")
        draw.ellipse((380, 310, 820, 500), fill="#D94C42")
        draw.text((370, 585), "酸甜有汁，鸡蛋松软，拌饭拌面都合适", font=label_font, fill=palette[1])

    img.save(path)


def build_images() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    images = {
        "chapter1": ("厨房里最稳的一道家常菜", "番茄和鸡蛋看似简单，真正好吃靠的是顺序与火候。", ("#FEF7F2", "#9E332A", "#845A54"), "plated"),
        "chapter2": ("食材选择与准备", "选对番茄、打好蛋液，成功已经完成一半。", ("#FFF8F0", "#9E332A", "#8A6056"), "ingredients"),
        "chapter3": ("核心技法拆解", "锅气、蛋香、番茄汁感，决定这道菜的层次。", ("#FFF7ED", "#8E3A2C", "#7F6053"), "heat"),
        "chapter4": ("分步操作路径", "把节奏拆成四拍，新手也能稳稳出锅。", ("#FFF9F2", "#A24833", "#876353"), "steps"),
        "chapter5": ("常见失误与修正", "出现问题不要慌，大多数都能及时救回来。", ("#FFF8F4", "#A44734", "#85645C"), "fix"),
        "chapter6": ("营养搭配与变化", "一盘菜也能吃得平衡，还能做出家里喜欢的版本。", ("#F8FBF6", "#5E7C4E", "#6D7D68"), "nutrition"),
        "chapter7": ("成菜效果参考", "颜色要亮、汁水要润、蛋块要松。", ("#FFF7F2", "#9D382C", "#8B5D4D"), "plated"),
        "prep": ("切配细节", "番茄切块不切碎，鸡蛋要先打出空气感。", ("#FFF8F3", "#964234", "#83655A"), "knife"),
    }
    result = {}
    for key, (title, subtitle, palette, mode) in images.items():
        path = ASSET_DIR / f"{key}.png"
        generate_illustration(path, title, subtitle, palette, mode)
        result[key] = path
    return result


def add_picture_with_caption(doc: Document, image_path: Path, caption: str, *, width=Cm(15.5)):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=width)
    add_caption(doc, caption)


def chapter_intro(doc: Document, heading: str, image_path: Path, caption: str):
    add_heading(doc, heading, level=1)
    add_picture_with_caption(doc, image_path, caption)


def create_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    image_map = build_images()

    doc = Document()
    section = doc.sections[0]
    set_section_layout(section)
    add_header_footer(section, "第一章 引言")

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_BODY
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    normal.font.size = Pt(12)

    title_p = doc.add_paragraph()
    style_paragraph(title_p, alignment=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.15, before=0, after=6)
    title_run = title_p.add_run("西红柿炒鸡蛋的最佳做法——从入门到精通的家常美味指南")
    set_east_asia_font(title_run, FONT_UI, 22, bold=True, color=TITLE_COLOR)

    subtitle_p = doc.add_paragraph()
    style_paragraph(subtitle_p, alignment=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.15, before=0, after=10)
    subtitle_run = subtitle_p.add_run("掌握火候与调味的黄金法则，做出餐厅级口感")
    set_east_asia_font(subtitle_run, FONT_BODY, 12, color=MUTED_COLOR)

    meta_p = doc.add_paragraph()
    style_paragraph(meta_p, alignment=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.15, before=0, after=10)
    meta_run = meta_p.add_run("文档类型：实用烹饪指南 / 食谱教程    目标读者：家常菜爱好者、烹饪新手、家庭厨师")
    set_east_asia_font(meta_run, FONT_BODY, 11, color=MUTED_COLOR)

    chapter_intro(doc, "第一章 引言", image_map["chapter1"], "图 1  理想的成菜状态应该是色泽鲜亮、汁水适中、蛋块松软。")
    add_body_paragraph(
        doc,
        "很多人第一次学做饭，都会从西红柿炒鸡蛋开始。它看起来家常，做法也不复杂，却是一道非常考验基本功的菜：番茄切太碎，锅里容易发水；鸡蛋炒得太老，整盘菜就会显得粗糙；盐糖比例不稳，酸甜就会互相打架。也正因为如此，这道菜特别适合拿来练习火候、顺序和调味的分寸感。",
    )
    add_body_paragraph(
        doc,
        "真正好吃的西红柿炒鸡蛋，不是把鸡蛋和番茄一起丢进锅里翻一翻，而是让两种食材各自发挥长处：鸡蛋负责香和软，番茄负责鲜和润，最后再在锅里完成一次短暂而精准的融合。只要掌握几个关键动作，新手也能做出颜色明快、口感有层次、拌饭拌面都很出彩的效果。",
    )

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_layout(section)
    add_header_footer(section, "第二章 食材选择与准备")
    chapter_intro(doc, "第二章 食材选择与准备", image_map["chapter2"], "图 2  食材选得准，后面的火候和调味才有发挥空间。")
    add_body_paragraph(
        doc,
        "番茄建议选成熟度高、用手轻按略有回弹的品种，表皮发亮、颜色红润、蒂部没有青硬感为佳。过生的番茄酸味冲、果肉紧，炒半天也不容易出汁；过熟到发绵的番茄虽然出汁快，但下锅后容易碎得一塌糊涂。家常做法里，普通圆番茄就很好用，如果能买到汁水更足的沙瓤番茄，成菜会更有“汤感”。",
    )
    add_body_paragraph(
        doc,
        "鸡蛋以新鲜为第一标准，蛋壳表面完整、无裂纹即可。三枚到四枚鸡蛋适合搭配中等大小番茄两到三个。打蛋时不要只是把蛋黄蛋清打散，而要顺着一个方向多打几十下，让蛋液里带进适量空气，这样下锅后更蓬松。此时加入少许盐和几滴料酒，底味和去腥都能一步到位。如果家里喜欢更嫩的口感，也可以加一小勺温水，让蛋液受热后更柔润。",
    )
    add_picture_with_caption(doc, image_map["prep"], "图 3  番茄切块更利于保留汁水，蛋液打到略起泡更容易蓬松。", width=Cm(14.8))
    add_body_paragraph(
        doc,
        "切配也有讲究。番茄去蒂后切成大小接近的滚刀块或月牙块，不必切得太小，否则很快炒烂；如果介意番茄皮影响口感，可以在顶部划十字后用热水烫十几秒，轻松去皮。葱花不是必须，但少量葱白能帮助提香。想让颜色更好看，可以把蛋液和番茄分开放置，临下锅前再确认调料是否备齐，这样操作时节奏不会乱。",
    )

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_layout(section)
    add_header_footer(section, "第三章 核心技法详解")
    chapter_intro(doc, "第三章 核心技法详解", image_map["chapter3"], "图 4  热锅宽油与分锅处理，是口感拉开差距的核心。")
    add_body_paragraph(
        doc,
        "第一项核心技法是“分炒再合”。鸡蛋和番茄熟成速度不同，若从头一起炒，鸡蛋为了等番茄出汁会越炒越老，番茄则会在鸡蛋吸附下失去清爽感。正确做法是先把鸡蛋炒到七八成熟盛出，再单独把番茄炒软、炒出红亮汁水，最后回锅合炒。这样鸡蛋保留松软感，番茄也保留鲜亮和存在感。",
    )
    add_body_paragraph(
        doc,
        "第二项核心技法是“让番茄自己出汁”。很多人习惯一开始就倒很多水，结果成菜像汤，不像菜。番茄本身水分充足，只要油温合适、锅里有一点盐帮助渗透，果肉自然会慢慢塌软并释放汁液。此时用锅铲轻压边缘，而不是不停乱翻，会更容易形成带一点浓度的自然酱汁。只有在番茄偏生、汁明显不足时，才补一两勺清水或番茄汁。",
    )
    add_body_paragraph(
        doc,
        "第三项关键是掌握鸡蛋成熟的“窗口期”。锅热后入油，油微微起纹就能下蛋液。蛋液入锅后先不要急着狂搅，等边缘轻轻鼓起，再从外向内推，形成大块柔软的蛋花。看到表面还有少许湿润感时就盛出，因为回锅时它还会再熟一次。这个阶段宁愿略嫩，也不要一次炒到全熟，这就是成菜嫩滑的秘诀。",
    )
    add_body_paragraph(
        doc,
        "最后是调味节奏。西红柿炒鸡蛋的酸甜鲜咸要平衡，最怕一种味道压住其他味道。最稳妥的方法是：蛋液里先放底盐，番茄炒软后放主盐和少量糖，回锅后再尝一口补最后的味。白糖的作用主要是托住番茄香气、修饰酸感，不是把菜做成甜口；如果番茄本身已经很甜，糖量可以明显减少。懂得分层调味，这道菜就不会“只有番茄味”或者“只剩鸡蛋香”。",
    )
    add_recipe_table(doc)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_layout(section)
    add_header_footer(section, "第四章 分步操作指南")
    chapter_intro(doc, "第四章 分步操作指南", image_map["chapter4"], "图 5  按照固定节奏推进，操作会非常稳。")
    add_step(doc, "准备食材", "3 枚鸡蛋配 2 到 3 个中等番茄最常见，番茄切块，鸡蛋打入碗中，加入少许盐和几滴料酒，顺着一个方向搅打到表面略起泡。")
    add_step(doc, "先炒鸡蛋", "锅烧热后倒入略多一点的油，轻轻晃锅让油铺开，蛋液下锅后等边缘鼓起，再用锅铲从外往里推，形成大块蛋花，八成熟时立刻盛出。")
    add_step(doc, "再炒番茄", "锅里补少量油，下葱白略煸香后放入番茄，中火翻几下，撒一点盐帮助出汁，看到番茄边缘变软后可轻压几下，让汁水更快释放。")
    add_step(doc, "建立味型", "番茄炒出红亮汁水后，加入白糖和少量清水，尝一下酸度。如果番茄偏酸，就稍多一点糖；如果番茄本身香甜，只需少量提鲜即可。")
    add_step(doc, "回锅合炒", "把鸡蛋重新倒回锅中，转大火快速翻匀 10 到 15 秒，让蛋块均匀裹上番茄汁。这个过程不要久煮，目的是让味道结合，而不是把鸡蛋炖老。")
    add_step(doc, "收汁出锅", "最后看锅里状态决定是否收汁：喜欢拌饭可留一点红亮汤汁，喜欢干香一点则大火再收十几秒。出锅前撒少量葱花，颜色和香气都会更完整。")
    add_body_paragraph(
        doc,
        "如果是第一次做，建议整道菜从头到尾都用中大火完成，但每一个动作都要短而利落。火太小，鸡蛋不香、番茄不亮；火太大又手忙脚乱，就容易糊边。记住一个简单原则：鸡蛋靠热油快速定型，番茄靠中火慢慢出汁，合炒靠大火迅速收尾。",
    )

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_layout(section)
    add_header_footer(section, "第五章 常见问题与解决方案")
    chapter_intro(doc, "第五章 常见问题与解决方案", image_map["chapter5"], "图 6  常见失误基本都能通过顺序、火候和调味修回来。")
    add_body_paragraph(
        doc,
        "西红柿炒鸡蛋最常见的失败，不是调料记错，而是节奏乱了。鸡蛋老、番茄散、汤汁稀、味道淡，背后往往都是同一类问题：锅不够热、动作不够干脆，或者把所有调味都堆到最后。好消息是，这些问题大都可复现、也可修正。只要知道每个现象对应哪一步失控，下次就能明显进步。",
    )
    add_problem_table(doc)
    add_body_paragraph(
        doc,
        "还有两个小细节经常被忽略。第一，鸡蛋盛出后不要放太久，最好番茄一出汁就回锅，否则温差太大、融合感会弱。第二，出锅前务必尝味，哪怕只是尝一小口汤汁，也能及时发现盐糖是否需要最后微调。家常菜的稳定感，往往就来自这一步。 ",
    )

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_layout(section)
    add_header_footer(section, "第六章 营养搭配与变化")
    chapter_intro(doc, "第六章 营养搭配与变化", image_map["chapter6"], "图 7  一荤一素一主食的搭配，能让这道菜更完整。")
    add_body_paragraph(
        doc,
        "从营养角度看，番茄提供维生素 C、番茄红素和清爽酸味，鸡蛋则补充优质蛋白和一定脂溶性营养，两者同炒不仅味道合拍，也很适合日常家庭餐桌。若想让吸收更好，适量油脂是必要的，因为番茄红素在有油环境下更容易释放和利用。这也是为什么太“清淡”的版本往往香气不足、营养表现也不理想。",
    )
    add_body_paragraph(
        doc,
        "在搭配上，这道菜最适合配米饭、面条、馒头等主食；如果想让一餐更均衡，可以再加一道清炒绿叶菜或凉拌黄瓜。家里有老人和孩子时，可把番茄切得稍小、鸡蛋炒得更嫩，整体会更易咀嚼。喜欢丰富口感的人，还可以加入木耳、虾仁或少量青椒，但原则是不抢主味，始终让番茄与鸡蛋做主角。",
    )
    add_bullet(doc, "偏酸爽版：选择成熟但酸感更明显的番茄，糖减到最低，适合夏天开胃。")
    add_bullet(doc, "偏浓香版：鸡蛋略煎出金边，番茄多收一点汁，适合拌饭。")
    add_bullet(doc, "轻油家常版：总油量稍减，但仍保留“先炒蛋、后炒番茄、最后合炒”的顺序。")

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_layout(section)
    add_header_footer(section, "第七章 结语")
    chapter_intro(doc, "第七章 结语", image_map["chapter7"], "图 8  理想的西红柿炒鸡蛋，应当同时具备香、鲜、润、亮四个特点。")
    add_body_paragraph(
        doc,
        "一道看似朴素的西红柿炒鸡蛋，其实把家常烹饪里最重要的几个基本功都串了起来：选材、预处理、火候、调味、节奏。把这道菜做好，不只是学会一个菜谱，更是在厨房里建立“先想清楚，再下锅”的做菜习惯。今后无论做青椒炒蛋、木须肉还是番茄牛腩，这种思路都能继续用得上。",
    )
    add_body_paragraph(
        doc,
        "下次做这道菜时，不妨只盯住三个目标：鸡蛋要嫩，番茄要出汁，味道要平衡。做到这三点，这盘家常菜就会从“能吃”变成“想再来一碗饭”。当你能稳定做出这一盘酸甜鲜香、颜色漂亮的西红柿炒鸡蛋时，厨艺已经悄悄往前迈了一大步。",
    )

    doc.core_properties.title = "西红柿炒鸡蛋的最佳做法——从入门到精通的家常美味指南"
    doc.core_properties.subject = "实用烹饪指南"
    doc.core_properties.author = "Codex"
    doc.core_properties.comments = "包含章节图示、调味比例表、问题速查表与页眉页脚。"
    doc.save(DOCX_PATH)


if __name__ == "__main__":
    create_document()
