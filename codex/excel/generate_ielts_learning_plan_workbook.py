from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


OUTPUT_PATH = Path("ielts-learning-plan.xlsx")

THIN = Side(style="thin", color="D9E2F3")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")


def style_headers(ws, row: int, cols: int) -> None:
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_table(ws, start_row: int, end_row: int, end_col: int) -> None:
    for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=1, max_col=end_col):
        for cell in row:
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def build_dashboard(ws) -> None:
    ws.title = "总览看板"
    ws.merge_cells("A1:L1")
    ws["A1"] = "雅思动态学习计划总览"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = HEADER_FILL
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    profile_labels = ["目标", "可用总时长", "注意力集中时长", "学习风格"]
    profile_values = ["雅思总分 [目标分]，单项不低于 [X]", "[X小时/周]", "[25分钟番茄钟 / 90分钟深度学习]", "[视觉型 / 听觉型 / 动手实践型 / 阅读型]"]
    for idx, (label, value) in enumerate(zip(profile_labels, profile_values), start=3):
        ws[f"A{idx}"] = label
        ws[f"A{idx}"].font = Font(bold=True)
        ws[f"A{idx}"].fill = SECTION_FILL
        ws[f"B{idx}"] = value
        ws[f"A{idx}"].border = ws[f"B{idx}"].border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    kpi_labels = [("J3", "总进度"), ("J4", "进度条"), ("J5", "已完成天数"), ("J6", "精通率")]
    for cell, label in kpi_labels:
        ws[cell] = label
        ws[cell].font = Font(bold=True)
        ws[cell].fill = SECTION_FILL
        ws[cell].border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    ws["K3"] = '=IFERROR((COUNTIF(\'周计划明细\'!$H$2:$H$200,"25%")*0.25+COUNTIF(\'周计划明细\'!$H$2:$H$200,"50%")*0.5+COUNTIF(\'周计划明细\'!$H$2:$H$200,"75%")*0.75+COUNTIF(\'周计划明细\'!$H$2:$H$200,"100%"))/MAX(1,COUNTA(\'周计划明细\'!$A$2:$A$200)),0)'
    ws["K4"] = '=REPT("|",ROUND(K3*20,0))&REPT(".",20-ROUND(K3*20,0))'
    ws["K5"] = '=COUNTIF(\'周计划明细\'!$H$2:$H$200,"100%")'
    ws["K6"] = '=IFERROR(COUNTIF(\'知识点追踪\'!$D$2:$D$200,"精通")/MAX(1,COUNTA(\'知识点追踪\'!$B$2:$B$200)),0)'
    ws["K3"].number_format = "0%"
    ws["K6"].number_format = "0%"

    headers = ["阶段", "时间节点", "核心模块", "里程碑", "验收标准", "当前状态", "阶段进度"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=8, column=col, value=header)
    style_headers(ws, 8, len(headers))

    rows = [
        ["阶段1 诊断与打底", "第1-2周", "词汇语法 / 听阅基础", "完成入门诊断与基础搭建", "完成1次模考并建立错题本，词汇首轮 300-500 词", "进行中", '=IF(F9="已完成",1,IF(F9="进行中",0.5,0))'],
        ["阶段2 单项强化", "第3-6周", "听力 / 阅读 / 写作 / 口语", "完成高频题型分项训练", "听阅正确率稳定提升，写作完成大小作文框架，口语完成题库首轮", "未开始", '=IF(F10="已完成",1,IF(F10="进行中",0.5,0))'],
        ["阶段3 冲刺输出", "第7-8周", "套题模考 / 复盘", "完成全真套题与查漏补缺", "模考达到目标分数区间，弱项得到补强", "未开始", '=IF(F11="已完成",1,IF(F11="进行中",0.5,0))'],
    ]
    for row_idx, row in enumerate(rows, start=9):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
        ws.cell(row=row_idx, column=7).number_format = "0%"
    style_table(ws, 9, 11, 7)

    summary_headers = ["模块", "计划投入(h)", "占比"]
    for col, header in enumerate(summary_headers, start=9):
        ws.cell(row=8, column=col, value=header)
    style_headers(ws, 8, 11)

    summary_modules = ["听力", "阅读", "写作", "口语", "词汇语法"]
    for row_idx, module in enumerate(summary_modules, start=9):
        ws.cell(row=row_idx, column=9, value=module)
        ws.cell(row=row_idx, column=10, value=f'=SUMIF(统计源!$C$2:$C$800,I{row_idx},统计源!$D$2:$D$800)')
        ws.cell(row=row_idx, column=11, value=f'=IFERROR(J{row_idx}/SUM($J$9:$J$13),0)')
        ws.cell(row=row_idx, column=11).number_format = "0%"
    style_table(ws, 9, 13, 11)

    pie = PieChart()
    pie.title = "各模块投入时间占比"
    data = Reference(ws, min_col=10, min_row=8, max_row=13)
    labels = Reference(ws, min_col=9, min_row=9, max_row=13)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.height = 7
    pie.width = 9
    ws.add_chart(pie, "I15")

    log_headers = ["调整日期", "变更内容", "变更原因", "影响范围", "应对措施"]
    for col, header in enumerate(log_headers, start=1):
        ws.cell(row=15, column=col, value=header)
    style_headers(ws, 15, len(log_headers))

    log_rows = [
        [date(2026, 7, 27), "阅读长难句改到上午", "晚间处理复杂材料效率低", "阅读训练", "把错题复盘移到晚上"],
        [date(2026, 8, 3), "增加 1 次口语录音回放", "流利度提升慢", "口语训练", "每周固定一次自评纠音"],
        ["", "", "", "", ""],
    ]
    for row_idx, row in enumerate(log_rows, start=16):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    style_table(ws, 16, 18, 5)

    ws.freeze_panes = "A8"
    for col, width in {"A": 16, "B": 24, "C": 22, "D": 20, "E": 30, "F": 12, "G": 12, "I": 14, "J": 14, "K": 12}.items():
        ws.column_dimensions[col].width = width


def build_week_plan(ws) -> None:
    ws.title = "周计划明细"
    headers = ["周次", "日期", "星期", "上午内容", "下午内容", "晚上内容", "学习时长(自动求和)", "完成度", "实际用时", "差异分析"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)
    style_headers(ws, 1, len(headers))

    rows = [
        [1, date(2026, 7, 20), '=TEXT(B2,"aaaa")', "阅读:长难句精读|1.5h", "听力:Section 2 精听|1.0h", "词汇语法:核心词复盘|0.5h", '=SUM(IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(D2,"|"),"h")),0),IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(E2,"|"),"h")),0),IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(F2,"|"),"h")),0))', "75%", 2.6, '=IF(OR(G2="",I2=""),"待填写",IF(I2>G2,"超时 "&TEXT(I2-G2,"0.0")&"h",IF(I2<G2,"节省 "&TEXT(G2-I2,"0.0")&"h","符合预期")))'],
        [1, date(2026, 7, 21), '=TEXT(B3,"aaaa")', "写作:Task 1 框架训练|1.5h", "阅读:判断题专项|1.0h", "口语:Part 2 录音复述|0.5h", '=SUM(IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(D3,"|"),"h")),0),IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(E3,"|"),"h")),0),IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(F3,"|"),"h")),0))', "100%", 3.0, '=IF(OR(G3="",I3=""),"待填写",IF(I3>G3,"超时 "&TEXT(I3-G3,"0.0")&"h",IF(I3<G3,"节省 "&TEXT(G3-I3,"0.0")&"h","符合预期")))'],
        [1, date(2026, 7, 22), '=TEXT(B4,"aaaa")', "听力:地图题专项|1.0h", "写作:Task 2 论证段练习|1.5h", "口语:Part 1 高频话题|0.5h", '=SUM(IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(D4,"|"),"h")),0),IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(E4,"|"),"h")),0),IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(F4,"|"),"h")),0))', "50%", 2.4, '=IF(OR(G4="",I4=""),"待填写",IF(I4>G4,"超时 "&TEXT(I4-G4,"0.0")&"h",IF(I4<G4,"节省 "&TEXT(G4-I4,"0.0")&"h","符合预期")))'],
    ]
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    style_table(ws, 2, 200, len(headers))
    dv = DataValidation(type="list", formula1='"0%,25%,50%,75%,100%"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("H2:H200")

    ws.conditional_formatting.add("A2:J200", FormulaRule(formula=['$H2="100%"'], stopIfTrue=False, fill=GREEN_FILL))
    ws.conditional_formatting.add("A2:J200", FormulaRule(formula=['=AND($B2<TODAY(),$H2<>"100%",$A2<>"")'], stopIfTrue=False, fill=RED_FILL))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:J200"
    for col, width in {"A": 8, "B": 12, "C": 10, "D": 24, "E": 24, "F": 24, "G": 14, "H": 12, "I": 10, "J": 18}.items():
        ws.column_dimensions[col].width = width


def build_knowledge(ws) -> None:
    ws.title = "知识点追踪"
    headers = ["模块", "知识点", "难度等级(1-5星)", "掌握程度", "首次学习日期", "复习日期1", "复习日期2", "复习日期3", "关联知识点", "复习日期4", "复习日期5", "复习提醒"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)
    style_headers(ws, 1, len(headers))

    rows = [
        ["听力", "Section 2 场景词", "★★★☆☆", "熟悉", date(2026, 7, 20), "=E2+1", "=E2+2", "=E2+4", "同义替换", "=E2+7", "=E2+15", '=IF(E2="","",IF(TODAY()>K2,"应进行第5次复习",IF(TODAY()>J2,"应进行第4次复习",IF(TODAY()>H2,"应进行第3次复习",IF(TODAY()>G2,"应进行第2次复习",IF(TODAY()>F2,"应进行第1次复习","未到复习日"))))))'],
        ["阅读", "判断题定位", "★★★★☆", "生疏", date(2026, 7, 21), "=E3+1", "=E3+2", "=E3+4", "段落主旨题", "=E3+7", "=E3+15", '=IF(E3="","",IF(TODAY()>K3,"应进行第5次复习",IF(TODAY()>J3,"应进行第4次复习",IF(TODAY()>H3,"应进行第3次复习",IF(TODAY()>G3,"应进行第2次复习",IF(TODAY()>F3,"应进行第1次复习","未到复习日"))))))'],
        ["写作", "Task 2 让步转折段", "★★★★★", "生疏", date(2026, 7, 22), "=E4+1", "=E4+2", "=E4+4", "论证展开", "=E4+7", "=E4+15", '=IF(E4="","",IF(TODAY()>K4,"应进行第5次复习",IF(TODAY()>J4,"应进行第4次复习",IF(TODAY()>H4,"应进行第3次复习",IF(TODAY()>G4,"应进行第2次复习",IF(TODAY()>F4,"应进行第1次复习","未到复习日"))))))'],
    ]
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    style_table(ws, 2, 200, 12)
    mastery_dv = DataValidation(type="list", formula1='"生疏,熟悉,精通"', allow_blank=True)
    ws.add_data_validation(mastery_dv)
    mastery_dv.add("D2:D200")
    ws.conditional_formatting.add("A2:L200", FormulaRule(formula=['=AND($L2<>"未到复习日",$D2<>"精通",$A2<>"")'], stopIfTrue=False, fill=YELLOW_FILL))

    for col, width in {"A": 12, "B": 22, "C": 16, "D": 12, "E": 14, "F": 12, "G": 12, "H": 12, "I": 16, "J": 12, "K": 12, "L": 16}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.column_dimensions["J"].hidden = True
    ws.column_dimensions["K"].hidden = True
    ws.column_dimensions["L"].hidden = True


def build_resources(ws) -> None:
    ws.title = "资源清单"
    headers = ["资源名称", "类型(书/视频/题库/文档)", "对应模块", "预估耗时", "已用时长", "完成状态", "链接/位置", "质量评分"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)
    style_headers(ws, 1, len(headers))

    rows = [
        ["Cambridge IELTS 真题集", "题库", "听力", 12, 4, "进行中", "纸质/电子版", 4.9],
        ["雅思阅读长难句精讲", "文档", "阅读", 8, 2, "进行中", "本地资料夹", 4.6],
        ["Simon 写作范文库", "文档", "写作", 6, 1, "未开始", "网盘/笔记", 4.7],
    ]
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    style_table(ws, 2, 200, len(headers))
    type_dv = DataValidation(type="list", formula1='"书,视频,题库,文档"', allow_blank=True)
    status_dv = DataValidation(type="list", formula1='"未开始,进行中,已完成"', allow_blank=True)
    ws.add_data_validation(type_dv)
    ws.add_data_validation(status_dv)
    type_dv.add("B2:B200")
    status_dv.add("F2:F200")

    for col, width in {"A": 24, "B": 20, "C": 12, "D": 12, "E": 12, "F": 12, "G": 20, "H": 10}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"


def build_stats(ws) -> None:
    ws.title = "统计源"
    ws.append(["来源行", "时段", "模块", "计划时长"])
    style_headers(ws, 1, 4)
    current_row = 2
    for plan_row in range(2, 201):
        for period_name, col in [("上午", "D"), ("下午", "E"), ("晚上", "F")]:
            ws.cell(row=current_row, column=1, value=plan_row)
            ws.cell(row=current_row, column=2, value=period_name)
            ws.cell(row=current_row, column=3, value=f'=IFERROR(TEXTBEFORE(\'周计划明细\'!{col}{plan_row},":"),"")')
            ws.cell(row=current_row, column=4, value=f'=IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(\'周计划明细\'!{col}{plan_row},"|"),"h")),0)')
            current_row += 1
    style_table(ws, 2, current_row - 1, 4)
    ws.sheet_state = "hidden"


def apply_formats(wb: Workbook) -> None:
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"
                elif isinstance(cell.value, float):
                    cell.number_format = "0.0"


def main() -> None:
    wb = Workbook()
    build_dashboard(wb.active)
    build_week_plan(wb.create_sheet())
    build_knowledge(wb.create_sheet())
    build_resources(wb.create_sheet())
    build_stats(wb.create_sheet())
    apply_formats(wb)
    wb.save(OUTPUT_PATH)
    print(OUTPUT_PATH.resolve())


if __name__ == "__main__":
    main()
