from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


OUTPUT_PATH = Path("动态学习计划模板.xlsx")


THIN = Side(style="thin", color="D9E2F3")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")


def style_headers(ws, row: int, cols: int) -> None:
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_table(ws, start_row: int, end_row: int, end_col: int) -> None:
    for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=1, max_col=end_col):
        for cell in row:
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def build_dashboard(ws) -> None:
    ws.title = "总览看板"
    ws.merge_cells("A1:L1")
    ws["A1"] = "动态学习计划总览看板"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = HEADER_FILL
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    profile_labels = ["目标", "可用总时长", "注意力集中时长", "学习风格"]
    profile_values = ["[具体目标]", "[X小时/周]", "[25分钟番茄钟 / 90分钟深度学习]", "[视觉型 / 听觉型 / 动手实践型 / 阅读型]"]
    for idx, (label, value) in enumerate(zip(profile_labels, profile_values), start=3):
        ws[f"A{idx}"] = label
        ws[f"A{idx}"].font = Font(bold=True)
        ws[f"A{idx}"].fill = SECTION_FILL
        ws[f"B{idx}"] = value
        ws[f"A{idx}"].border = ws[f"B{idx}"].border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    kpi_labels = [("J3", "总进度"), ("J4", "进度条"), ("J5", "已完成天数"), ("J6", "知识点精通率")]
    for cell, label in kpi_labels:
        ws[cell] = label
        ws[cell].font = Font(bold=True)
        ws[cell].fill = SECTION_FILL
        ws[cell].border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    ws["K3"] = '=IFERROR((COUNTIF(\'周计划明细\'!$H$2:$H$200,"25%")*0.25+COUNTIF(\'周计划明细\'!$H$2:$H$200,"50%")*0.5+COUNTIF(\'周计划明细\'!$H$2:$H$200,"75%")*0.75+COUNTIF(\'周计划明细\'!$H$2:$H$200,"100%"))/MAX(1,COUNTA(\'周计划明细\'!$A$2:$A$200)),0)'
    ws["K4"] = '=REPT("|",ROUND(K3*20,0))&REPT(".",20-ROUND(K3*20,0))'
    ws["K5"] = '=COUNTIF(\'周计划明细\'!$H$2:$H$200,"100%")'
    ws["K6"] = '=IFERROR(COUNTIF(\'知识点追踪\'!$D$2:$D$200,"精通")/MAX(1,COUNTA(\'知识点追踪\'!$B$2:$B$200)),0)'
    ws["K3"].number_format = "0%"
    ws["K6"].number_format = "0%"

    roadmap_headers = ["阶段", "时间节点", "核心模块", "里程碑", "验收标准", "当前状态", "阶段进度"]
    for col, header in enumerate(roadmap_headers, start=1):
        ws.cell(row=8, column=col, value=header)
    style_headers(ws, 8, len(roadmap_headers))

    roadmap_rows = [
        ["阶段1 启动", "第1-2周", "基础理论", "完成知识框架搭建", "能复述核心概念并完成基础题 80%", "进行中", '=IF(F9="已完成",1,IF(F9="进行中",0.5,0))'],
        ["阶段2 强化", "第3-6周", "核心技能", "完成重点模块训练", "能独立完成综合练习并输出错题清单", "未开始", '=IF(F10="已完成",1,IF(F10="进行中",0.5,0))'],
        ["阶段3 输出", "第7-8周", "实战应用", "完成项目/模拟考/复盘", "达到目标分数或完成项目交付", "未开始", '=IF(F11="已完成",1,IF(F11="进行中",0.5,0))'],
    ]
    for row_idx, row in enumerate(roadmap_rows, start=9):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
        ws.cell(row=row_idx, column=7).number_format = "0%"

    style_table(ws, 9, 11, 7)

    summary_headers = ["模块", "计划投入(h)", "占比"]
    for col, header in enumerate(summary_headers, start=9):
        ws.cell(row=8, column=col, value=header)
    style_headers(ws, 8, 11)

    summary_modules = ["基础理论", "核心技能", "实战应用", "复盘输出", "机动调整"]
    for row_idx, module in enumerate(summary_modules, start=9):
        ws.cell(row=row_idx, column=9, value=module)
        ws.cell(row=row_idx, column=10, value=f'=SUMIF(统计源!$C$2:$C$800,I{row_idx},统计源!$D$2:$D$800)')
        ws.cell(row=row_idx, column=11, value=f'=IFERROR(J{row_idx}/SUM($J$9:$J$13),0)')
        ws.cell(row=row_idx, column=11).number_format = "0%"
    style_table(ws, 9, 13, 11)

    pie = PieChart()
    pie.title = "各模块投入时间占比"
    data = Reference(ws, min_col=10, min_row=8, max_row=13)
    labels = Reference(ws, min_col=9, min_row=9, max_row=13)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.height = 7
    pie.width = 9
    ws.add_chart(pie, "I15")

    log_headers = ["调整日期", "变更内容", "变更原因", "影响范围", "应对措施"]
    for col, header in enumerate(log_headers, start=1):
        ws.cell(row=15, column=col, value=header)
    style_headers(ws, 15, len(log_headers))

    log_rows = [
        [date(2026, 7, 27), "晚间学习改为错题复盘", "连续两天注意力下降", "晚间时段", "将重内容前移到上午"],
        [date(2026, 8, 3), "周末增加 2 小时机动时段", "阶段测验暴露薄弱点", "周末计划", "补强化练习与复盘"],
        ["", "", "", "", ""],
    ]
    for row_idx, row in enumerate(log_rows, start=16):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    style_table(ws, 16, 18, 5)

    ws.freeze_panes = "A8"
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 28
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["I"].width = 14
    ws.column_dimensions["J"].width = 14
    ws.column_dimensions["K"].width = 12


def build_week_plan(ws) -> None:
    ws.title = "周计划明细"
    headers = ["周次", "日期", "星期", "上午内容", "下午内容", "晚上内容", "学习时长(自动求和)", "完成度", "实际用时", "差异分析"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)
    style_headers(ws, 1, len(headers))

    sample_rows = [
        [1, date(2026, 7, 20), '=TEXT(B2,"aaaa")', "基础理论:课程导学|1.5h", "核心技能:例题拆解|1.0h", "复盘输出:学习笔记|0.5h", '=SUM(IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(D2,"|"),"h")),0),IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(E2,"|"),"h")),0),IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(F2,"|"),"h")),0))', "75%", 1.8, '=IF(OR(G2="",I2=""),"待填写",IF(I2>G2,"超时 "&TEXT(I2-G2,"0.0")&"h",IF(I2<G2,"节省 "&TEXT(G2-I2,"0.0")&"h","符合预期")))'],
        [1, date(2026, 7, 21), '=TEXT(B3,"aaaa")', "核心技能:章节练习|2.0h", "实战应用:案例模仿|1.0h", "复盘输出:错题整理|0.5h", '=SUM(IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(D3,"|"),"h")),0),IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(E3,"|"),"h")),0),IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(F3,"|"),"h")),0))', "100%", 3.5, '=IF(OR(G3="",I3=""),"待填写",IF(I3>G3,"超时 "&TEXT(I3-G3,"0.0")&"h",IF(I3<G3,"节省 "&TEXT(G3-I3,"0.0")&"h","符合预期")))'],
        [1, date(2026, 7, 22), '=TEXT(B4,"aaaa")', "基础理论:知识回顾|1.0h", "核心技能:专项训练|1.5h", "机动调整:补漏复习|0.5h", '=SUM(IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(D4,"|"),"h")),0),IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(E4,"|"),"h")),0),IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(F4,"|"),"h")),0))', "50%", 2.2, '=IF(OR(G4="",I4=""),"待填写",IF(I4>G4,"超时 "&TEXT(I4-G4,"0.0")&"h",IF(I4<G4,"节省 "&TEXT(G4-I4,"0.0")&"h","符合预期")))'],
    ]
    for row_idx, row in enumerate(sample_rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    style_table(ws, 2, 200, len(headers))
    dv = DataValidation(type="list", formula1='"0%,25%,50%,75%,100%"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("H2:H200")

    green_rule = FormulaRule(formula=['$H2="100%"'], stopIfTrue=False, fill=GREEN_FILL)
    red_rule = FormulaRule(formula=['=AND($B2<TODAY(),$H2<>"100%",$A2<>"")'], stopIfTrue=False, fill=RED_FILL)
    ws.conditional_formatting.add("A2:J200", green_rule)
    ws.conditional_formatting.add("A2:J200", red_rule)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:J200"
    widths = {"A": 8, "B": 12, "C": 10, "D": 26, "E": 26, "F": 26, "G": 14, "H": 12, "I": 10, "J": 18}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def build_knowledge(ws) -> None:
    ws.title = "知识点追踪"
    headers = ["模块", "知识点", "难度等级(1-5星)", "掌握程度", "首次学习日期", "复习日期1", "复习日期2", "复习日期3", "关联知识点"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)
    style_headers(ws, 1, len(headers))

    ws["J1"] = "复习日期4"
    ws["K1"] = "复习日期5"
    ws["L1"] = "复习提醒"
    style_headers(ws, 1, 12)

    sample_rows = [
        ["基础理论", "函数极限", "★★★☆☆", "熟悉", date(2026, 7, 20), "=E2+1", "=E2+2", "=E2+4", "导数定义", "=E2+7", "=E2+15", '=IF(E2="","",IF(TODAY()>K2,"应进行第5次复习",IF(TODAY()>J2,"应进行第4次复习",IF(TODAY()>H2,"应进行第3次复习",IF(TODAY()>G2,"应进行第2次复习",IF(TODAY()>F2,"应进行第1次复习","未到复习日"))))))'],
        ["核心技能", "题型拆解", "★★★★☆", "生疏", date(2026, 7, 21), "=E3+1", "=E3+2", "=E3+4", "错题归因", "=E3+7", "=E3+15", '=IF(E3="","",IF(TODAY()>K3,"应进行第5次复习",IF(TODAY()>J3,"应进行第4次复习",IF(TODAY()>H3,"应进行第3次复习",IF(TODAY()>G3,"应进行第2次复习",IF(TODAY()>F3,"应进行第1次复习","未到复习日"))))))'],
        ["实战应用", "综合案例", "★★★★★", "生疏", date(2026, 7, 22), "=E4+1", "=E4+2", "=E4+4", "项目复盘", "=E4+7", "=E4+15", '=IF(E4="","",IF(TODAY()>K4,"应进行第5次复习",IF(TODAY()>J4,"应进行第4次复习",IF(TODAY()>H4,"应进行第3次复习",IF(TODAY()>G4,"应进行第2次复习",IF(TODAY()>F4,"应进行第1次复习","未到复习日"))))))'],
    ]

    for row_idx, row in enumerate(sample_rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    style_table(ws, 2, 200, 12)

    mastery_dv = DataValidation(type="list", formula1='"生疏,熟悉,精通"', allow_blank=True)
    ws.add_data_validation(mastery_dv)
    mastery_dv.add("D2:D200")

    due_fill_rule = FormulaRule(formula=['=AND($L2<>"未到复习日",$D2<>"精通",$A2<>"")'], stopIfTrue=False, fill=YELLOW_FILL)
    ws.conditional_formatting.add("A2:L200", due_fill_rule)

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 12
    ws.column_dimensions["I"].width = 18
    ws.column_dimensions["J"].width = 12
    ws.column_dimensions["K"].width = 12
    ws.column_dimensions["L"].width = 16
    ws.freeze_panes = "A2"
    ws.column_dimensions["J"].hidden = True
    ws.column_dimensions["K"].hidden = True
    ws.column_dimensions["L"].hidden = True


def build_resources(ws) -> None:
    ws.title = "资源清单"
    headers = ["资源名称", "类型(书/视频/题库/文档)", "对应模块", "预估耗时", "已用时长", "完成状态", "链接/位置", "质量评分"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)
    style_headers(ws, 1, len(headers))

    sample_rows = [
        ["《基础课程讲义》", "书", "基础理论", 6, 2, "进行中", "本地/教材", 4.5],
        ["专项题训练营", "题库", "核心技能", 10, 4, "进行中", "在线题库", 4.8],
        ["案例拆解视频", "视频", "实战应用", 5, 0, "未开始", "网盘/课程平台", 4.2],
    ]
    for row_idx, row in enumerate(sample_rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    style_table(ws, 2, 200, len(headers))

    type_dv = DataValidation(type="list", formula1='"书,视频,题库,文档"', allow_blank=True)
    status_dv = DataValidation(type="list", formula1='"未开始,进行中,已完成"', allow_blank=True)
    ws.add_data_validation(type_dv)
    ws.add_data_validation(status_dv)
    type_dv.add("B2:B200")
    status_dv.add("F2:F200")

    widths = {"A": 22, "B": 20, "C": 14, "D": 12, "E": 12, "F": 12, "G": 22, "H": 10}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"


def build_stats(ws) -> None:
    ws.title = "统计源"
    ws.append(["来源行", "时段", "模块", "计划时长"])
    style_headers(ws, 1, 4)

    current_row = 2
    for plan_row in range(2, 201):
        for period_name, col in [("上午", "D"), ("下午", "E"), ("晚上", "F")]:
            ws.cell(row=current_row, column=1, value=plan_row)
            ws.cell(row=current_row, column=2, value=period_name)
            ws.cell(row=current_row, column=3, value=f'=IFERROR(TEXTBEFORE(\'周计划明细\'!{col}{plan_row},":"),"")')
            ws.cell(row=current_row, column=4, value=f'=IFERROR(VALUE(TEXTBEFORE(TEXTAFTER(\'周计划明细\'!{col}{plan_row},"|"),"h")),0)')
            current_row += 1

    style_table(ws, 2, current_row - 1, 4)
    ws.sheet_state = "hidden"


def apply_formats(wb: Workbook) -> None:
    date_format = "yyyy-mm-dd"
    decimal_format = "0.0"

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, date):
                    cell.number_format = date_format
                elif isinstance(cell.value, float):
                    cell.number_format = decimal_format


def main() -> None:
    wb = Workbook()
    dashboard = wb.active
    build_dashboard(dashboard)
    build_week_plan(wb.create_sheet())
    build_knowledge(wb.create_sheet())
    build_resources(wb.create_sheet())
    build_stats(wb.create_sheet())
    apply_formats(wb)
    wb.save(OUTPUT_PATH)
    print(OUTPUT_PATH.resolve())


if __name__ == "__main__":
    main()
